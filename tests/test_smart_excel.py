"""Regression for the user's real XLS sample and generic non-BOM spreadsheets."""
from __future__ import annotations

import io
import os
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from filedesk.engine import OPERATIONS, FileDeskError, find_libreoffice
from filedesk.server import app

REAL_SAMPLE = Path(os.environ.get('FILEDESK_QA_BOM','/nonexistent/filedesk-private-qa-bom.xls'))

@pytest.mark.skipif(not REAL_SAMPLE.is_file(), reason='private user sample not shipped with project')
def test_real_xls_to_editable_word(tmp_path):
    if not find_libreoffice():pytest.skip('LibreOffice absent')
    source=REAL_SAMPLE
    output=OPERATIONS['excel_to_word']([source],tmp_path,{})[0]
    doc=Document(output)
    assert len(doc.tables)>=3
    main=next(t for t in doc.tables if t.rows[0].cells[0].text=='序号' and any(c.text=='产品编号' for c in t.rows[0].cells))
    rows=[[c.text for c in row.cells] for row in main.rows]
    assert len(rows)==12 # header + 11 subparts
    assert any('W14007000011' in ' '.join(r) and '204' in ' '.join(r) for r in rows)
    assert any('W14014000457' in ' '.join(r) and '50' in ' '.join(r) for r in rows)
    assert all('启启' not in ' '.join(r) for r in rows)
    all_text=' '.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert 'WLQD20260929034' in all_text
    assert '2026-09-29' in all_text
    assert 'TX26044' in all_text

@pytest.mark.skipif(not REAL_SAMPLE.is_file(), reason='private user sample not shipped with project')
def test_real_xls_to_readable_pdf(tmp_path):
    if not find_libreoffice():pytest.skip('LibreOffice absent')
    output=OPERATIONS['excel_smart_pdf']([REAL_SAMPLE],tmp_path,{})[0]
    reader=PdfReader(output)
    assert len(reader.pages)==2
    text='\n'.join(page.extract_text() for page in reader.pages)
    assert '按钮' in text and '施耐德' in text and 'W14007000011' in text
    assert '204' in text and '通信电缆' in text


def test_generic_excel_complete_fields_and_api(tmp_path):
    sample=Path(__file__).parent/'fixtures'/'orders.xlsx'
    client=TestClient(app)
    with sample.open('rb') as f:
        response=client.post('/api/process',data={'operation':'excel_to_word','options':'{}'},files=[('files',('orders.xlsx',f,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))])
    assert response.status_code==200,response.text
    dest=tmp_path/'orders.docx';dest.write_bytes(response.content)
    doc=Document(dest)
    text='\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert all(v in text for v in ('甲客户','乙客户','丙客户','A001','A002','A003','B1','B2','B3','待处理','已完成'))
    assert doc.tables


def test_unsupported_workbook_type_rejected(tmp_path):
    with pytest.raises(FileDeskError,match='xls'):
        OPERATIONS['excel_to_word']([tmp_path/'presentation.pptx'],tmp_path,{})


def test_headerless_data_rejected_instead_of_lost():
    from filedesk.smart_excel import _sheet_model
    # Second column with no label must not silently vanish.
    with pytest.raises(FileDeskError, match='未命名列'):
        _sheet_model('异常表', [['产品','', '数量'],['键盘','信息',2]])
