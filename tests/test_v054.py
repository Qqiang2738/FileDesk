"""Regressions from user's real two-page BOM PDF and broken Mac PDF rendering.
Tests generate their own files; the user's business documents are never bundled.
"""
from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader
from filedesk.engine import office_to_pdf,pdf_to_word


def test_docx_chinese_to_pdf_uses_direct_text_and_stays_searchable(tmp_path):
    d=Document()
    d.add_heading('电气安装材料 · 物料清单',0)
    table=d.add_table(rows=3,cols=4)
    for i,row in enumerate([('产品编号','产品名称','品牌','数量'),
                            ('W14007000014','按钮','施耐德','17'),
                            ('W14008000060','指示灯','施耐德','204')]):
        for cell,value in zip(table.rows[i].cells,row): cell.text=value
    source=tmp_path/'source.docx';d.save(source)
    pdf=office_to_pdf([source],tmp_path,{})[0]
    assert pdf.name.endswith('_清晰排版.pdf')
    text=''.join(p.extract_text() for p in PdfReader(pdf).pages)
    assert '施耐德' in text and 'W14007000014' in text and '204' in text
    assert '电气安装材料' in text


def test_pdf_to_word_reconstructs_real_tables_from_our_pdf(tmp_path):
    from docx.shared import Pt
    d=Document()
    d.sections[0].page_width=Pt(842)
    d.sections[0].page_height=Pt(595)
    d.add_paragraph('物料清单')
    table=d.add_table(rows=3,cols=3)
    rows=[['产品编号','产品名称','数量'],['W14007000014','按钮','17'],['W14008000060','指示灯','204']]
    for row,content in zip(table.rows,rows):
        for cell,value in zip(row.cells,content):cell.text=value
    src=tmp_path/'src.docx';d.save(src)
    pdf=office_to_pdf([src],tmp_path,{'mode':'readable'})[0]
    docx=pdf_to_word([pdf],tmp_path,{'mode':'structured'})[0]
    result=Document(docx)
    joined=' '.join(c.text for t in result.tables for r in t.rows for c in r.cells)
    assert len(result.tables) >= 1
    assert all(code in joined for code in ('W14007000014','按钮','施耐德') if code != '施耐德')
    assert '204' in joined
    assert '物料清单' in ' '.join(p.text for p in result.paragraphs)


def test_structured_mode_on_plain_pdf_returns_editable_paragraphs(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    pdf=tmp_path/'plain.pdf'
    canvas=Canvas(str(pdf));canvas.drawString(100,730,'Simple test document');canvas.save()
    docx=pdf_to_word([pdf],tmp_path,{})[0]
    assert 'Simple test document' in ' '.join(p.text for p in Document(docx).paragraphs)


def test_no_deprecated_fitz_import_and_valid_build_version():
    root=Path(__file__).parents[1]
    for path in (root/'filedesk').glob('*.py'):
        assert not re.search(r'\bimport\s+fitz\b',path.read_text())
    html=(root/'web/index.html').read_text()
    js=(root/'web/app.js').read_text()
    assert 'name="filedesk-build" content="1.0.0"' in html
    assert "const BUILD = '1.0.0'" in js
    assert '智能表格（推荐，可编辑）' in js
