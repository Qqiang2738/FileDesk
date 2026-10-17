import io
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from fastapi.testclient import TestClient

from filedesk.engine import OPERATIONS, FileDeskError, parse_pages, find_libreoffice
from filedesk.server import app


def make_pdf(path: Path, pages: int = 3):
    c = canvas.Canvas(str(path))
    for n in range(pages):
        c.setFont('Helvetica', 18)
        c.drawString(60, 700, f'TEST PAGE {n+1}')
        c.showPage()
    c.save()
    return path


def test_pages_parse():
    assert parse_pages('1-3,5,3', 7) == [0, 1, 2, 4]
    assert parse_pages('', 4, require=False) == [0, 1, 2, 3]
    with pytest.raises(FileDeskError):
        parse_pages('1-999', 4)
    with pytest.raises(FileDeskError):
        parse_pages('0', 4)


def test_pdf_actions(tmp_path):
    a = make_pdf(tmp_path / 'one.pdf', 3)
    b = make_pdf(tmp_path / 'two.pdf', 2)
    merge = OPERATIONS['pdf_merge']([a, b], tmp_path, {})[0]
    assert len(PdfReader(merge).pages) == 5
    extracted = OPERATIONS['pdf_extract']([a], tmp_path, {'pages': '2-3'})[0]
    assert len(PdfReader(extracted).pages) == 2
    removed = OPERATIONS['pdf_remove']([a], tmp_path, {'pages': '2'})[0]
    assert len(PdfReader(removed).pages) == 2
    rotated = OPERATIONS['pdf_rotate']([a], tmp_path, {'angle': '90', 'pages': '2'})[0]
    assert PdfReader(rotated).pages[1].get('/Rotate') == 90
    splitted = OPERATIONS['pdf_split']([a], tmp_path, {})
    assert len(splitted) == 3
    assert all(len(PdfReader(p).pages) == 1 for p in splitted)
    compressed = OPERATIONS['pdf_compress']([a], tmp_path, {'mode': 'lossless'})[0]
    assert compressed.is_file() and compressed.stat().st_size <= a.stat().st_size


def test_image_actions(tmp_path):
    source = tmp_path / 'image.png'
    Image.new('RGBA', (180, 100), (50, 90, 100, 120)).save(source)
    outputs = OPERATIONS['image_convert']([source], tmp_path, {'format': 'jpg', 'quality': '85', 'max_side': '0'})
    with Image.open(outputs[0]) as result:
        assert result.format == 'JPEG' and result.mode == 'RGB'
    pdf = OPERATIONS['images_to_pdf']([source, source], tmp_path, {})[0]
    assert len(PdfReader(pdf).pages) == 2
    pictures = OPERATIONS['pdf_to_images']([pdf], tmp_path, {'format': 'png', 'dpi': '96'})
    assert len(pictures) == 2
    with Image.open(pictures[0]) as picture:
        assert picture.width > 100


def test_pdf_word_modes(tmp_path):
    pdf = make_pdf(tmp_path / 'input.pdf', 2)
    from docx import Document
    editable = OPERATIONS['pdf_to_word']([pdf], tmp_path, {'mode': 'text'})[0]
    assert 'TEST PAGE 1' in '\n'.join(p.text for p in Document(editable).paragraphs)
    visual = OPERATIONS['pdf_to_word']([pdf], tmp_path, {'mode': 'visual'})[0]
    assert len(Document(visual).inline_shapes) == 2


def test_office_convert(tmp_path):
    if not find_libreoffice():
        pytest.skip('LibreOffice not installed')
    from docx import Document
    source = tmp_path / 'sample.docx'
    doc = Document()
    doc.add_paragraph('OFFICE TO PDF TEST')
    doc.save(source)
    pdf = OPERATIONS['office_to_pdf']([source], tmp_path, {})[0]
    assert len(PdfReader(pdf).pages) >= 1


def test_local_http(tmp_path):
    client = TestClient(app)
    assert client.get('/').status_code == 200
    health = client.get('/api/health')
    assert health.json()['local_only'] is True
    pdf = make_pdf(tmp_path / 'one.pdf', 2)
    with open(pdf, 'rb') as f:
        resp = client.post('/api/process', data={'operation':'pdf_split', 'options':'{}'}, files=[('files', ('one.pdf', f, 'application/pdf'))])
    assert resp.status_code == 200, resp.text
    assert resp.headers.get('content-type') == 'application/zip'
    with zipfile.ZipFile(io.BytesIO(resp.content)) as z:
        assert len(z.namelist()) == 2
    bad = client.post('/api/process', data={'operation':'unknown', 'options':'{}'}, files=[('files', ('a.pdf', b'x', 'application/pdf'))])
    assert bad.status_code == 400


def test_origin_and_file_validation(tmp_path):
    client = TestClient(app)
    response = client.post('/api/process', headers={'Origin': 'http://evil.test'}, data={'operation':'pdf_merge','options':'{}'}, files=[('files', ('test.pdf', b'not-pdf', 'application/pdf'))])
    assert response.status_code == 403
    response = client.post('/api/process', data={'operation':'pdf_merge','options':'{}'}, files=[('files', ('wrong.txt', b'text', 'text/plain'))])
    assert response.status_code == 400
