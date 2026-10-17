"""Regression coverage for persistent local result history and PDF quality."""
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfReader

from filedesk.pdf_quality import verify_chinese_pdf
from filedesk.engine import FileDeskError, office_to_pdf, find_libreoffice
from filedesk.result_store import ResultStore
from filedesk.server import app


def test_history_api_retrieve_after_refresh_and_clear(tmp_path, monkeypatch):
    monkeypatch.setenv('FILEDESK_HISTORY_DIR', str(tmp_path / 'private-history'))
    src = tmp_path / 'result.txt'
    with TestClient(app) as client:
        # Real processing endpoint; upload a tiny image and re-encode it.
        from PIL import Image
        image = Image.new('RGB', (15, 15), '#226644')
        bytesio = io.BytesIO(); image.save(bytesio, 'PNG')
        response = client.post('/api/process',data={'operation':'image_convert','options':'{"format":"webp"}'},files=[('files',('example.png',bytesio.getvalue(),'image/png'))])
        assert response.status_code == 200, response.text
        rid = response.headers['X-FileDesk-Result-Id']
        assert len(rid) == 32
        history = client.get('/api/history').json()['items']
        assert len(history) == 1 and history[0]['id'] == rid
        assert not any(item.name == 'example.png' for item in (tmp_path/'private-history').iterdir())
        saved = client.get(f'/api/history/{rid}/download')
        assert saved.status_code == 200 and saved.content == response.content
        assert saved.headers['content-disposition']
        assert client.get('/api/history/..%2Fetc%2Fpasswd/download').status_code != 200
        assert client.delete('/api/history').json()['ok']
        assert client.get('/api/history').json()['items'] == []
        assert client.get(f'/api/history/{rid}/download').status_code == 404


def test_history_retention_limits(tmp_path):
    store = ResultStore(tmp_path)
    for i in range(23):
        file = tmp_path / f'f{i}.pdf'
        file.write_bytes(b'%PDF-1.4\n' + str(i).encode())
        store.add(file, 'pdf_merge', 1)
    rows = store.list()
    assert len(rows) == 20
    assert all(store.get(row['id']) for row in rows)
    assert all(not row['name'].startswith('f0.') for row in rows)
    store.clear()
    assert store.list() == []


def test_embedded_unknown_font_name_does_not_fail(tmp_path, monkeypatch):
    import filedesk.pdf_quality as quality
    expected='物料名称安装设备电气测试工程师清单'
    monkeypatch.setattr(quality,'pdf_diagnostics',lambda p:{'text':expected,'pages':1,'fonts':[{'name':'/ABCDEF+CoreTextFontSubset', 'embedded':True,'unicode':True}]})
    quality.verify_chinese_pdf(tmp_path/'dummy.pdf', expected)


def test_mangled_chinese_is_rejected_despite_embedded_fonts(tmp_path, monkeypatch):
    import filedesk.pdf_quality as quality
    monkeypatch.setattr(quality,'pdf_diagnostics',lambda p:{'text':'文'*60,'pages':1,'fonts':[{'name':'/CJKSubset','embedded':True,'unicode':True}]})
    with pytest.raises(FileDeskError,match='内容校验未通过'):
        quality.verify_chinese_pdf(tmp_path/'dummy.pdf','物料安装电气测试设备产品清单工程师编号')


def test_unembedded_font_is_rejected_even_with_correct_text(tmp_path, monkeypatch):
    import filedesk.pdf_quality as quality
    expected='产品编号材料安装电气工程师记录'
    monkeypatch.setattr(quality,'pdf_diagnostics',lambda p:{'text':expected,'pages':1,'fonts':[{'name':'/KnownFont','embedded':False,'unicode':True}]})
    with pytest.raises(FileDeskError,match='未能确认嵌入'):
        quality.verify_chinese_pdf(tmp_path/'dummy.pdf',expected)


@pytest.mark.skipif(not find_libreoffice(), reason='LibreOffice not installed')
def test_actual_chinese_word_pdf(tmp_path):
    doc = Document()
    doc.add_heading('电气安装材料清单', level=1)
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ('产品编号','产品名称','数量')):cell.text=value
    for cell, value in zip(table.rows[1].cells, ('AA001','点火测试设备','17')):cell.text=value
    src = tmp_path / '材料.docx'; doc.save(src)
    pdf = office_to_pdf([src],tmp_path,{})[0]
    verify_chinese_pdf(pdf, '电气安装材料清单产品编号产品名称数量点火测试设备')
    assert len(PdfReader(pdf).pages)>0
