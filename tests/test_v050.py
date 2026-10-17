"""Browser-only V0.5.0 regression tests. No user documents are distributed."""
from pathlib import Path
from tempfile import TemporaryDirectory
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
import unittest

from filedesk.xlsx_reader import read_workbook
from filedesk.smart_excel import excel_to_word, excel_smart_pdf
from filedesk.pdf_quality import verify_chinese_pdf
from filedesk.engine import find_libreoffice, office_to_pdf

FIXTURE=Path(__file__).parent/'fixtures/orders.xlsx'
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def sample_rows(path: Path, count: int = 620):
    # Modify a small synthetic XLSX fixture with stdlib OpenXML; avoids including personal files.
    with ZipFile(FIXTURE) as src:
        contents={n:src.read(n) for n in src.namelist()}
    name='xl/worksheets/sheet1.xml'
    root=ET.fromstring(contents[name]);data=root.find(f'{{{NS}}}sheetData')
    data.clear()
    headings=['产品编号','产品名称','产品型号','数量','单位']
    for ri in range(1,count+2):
        row=ET.SubElement(data,f'{{{NS}}}row',{'r':str(ri)})
        vals=headings if ri==1 else [f'W{ri:06d}','测试物料',f'型号-{ri}',str(ri),'个']
        for col,value in enumerate(vals):
            cell=ET.SubElement(row,f'{{{NS}}}c',{'r':f'{chr(65+col)}{ri}','t':'n' if col==3 and ri>1 else 'str'})
            ET.SubElement(cell,f'{{{NS}}}v').text=value
    contents[name]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    with ZipFile(path,'w',ZIP_DEFLATED) as dest:
        for name,body in contents.items():dest.writestr(name,body)


class LargeSheetsTest(unittest.TestCase):
    def test_620_rows_split_into_usable_word_parts(self):
        from docx import Document
        with TemporaryDirectory() as tmp:
            work=Path(tmp);source=work/'large.xlsx'
            sample_rows(source,620)
            workbook=read_workbook(source)
            self.assertEqual(len(workbook[0][1]),621)
            out=work/'result';out.mkdir()
            parts=excel_to_word([source],out,{})
            self.assertEqual(len(parts),2)
            all_values=[]
            for docx in parts:
                doc=Document(docx)
                all_values.extend(c.text for table in doc.tables for row in table.rows for c in row.cells)
            self.assertIn('W000002',all_values)
            self.assertIn('W000621',all_values)

    @unittest.skipUnless(find_libreoffice(), 'LibreOffice not installed')
    def test_word_pdf_chinese_glyphs(self):
        from docx import Document
        with TemporaryDirectory() as tmp:
            work=Path(tmp);src=work/'材料验收.docx'
            doc=Document();doc.add_paragraph('材料验收记录：电气安装、测试设备、物料清单。');doc.save(src)
            converted=office_to_pdf([src],work,{})
            self.assertEqual(len(converted),1)
            verify_chinese_pdf(converted[0], expected_text='材料验收记录：电气安装、测试设备、物料清单。')


if __name__=='__main__': unittest.main()
