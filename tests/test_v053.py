"""Regression: genuine CJK PDF export, not just an ever-looser font checker."""
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from filedesk.pdf_quality import convert_with_font_retries, verify_chinese_pdf
from filedesk.readable_pdf import readable_docx_to_pdf


def make_docx(path: Path):
    d=Document()
    d.add_heading('电气安装材料清单',0)
    d.add_paragraph('中文能够正常显示，原始资料需要完整。')
    table=d.add_table(rows=1,cols=4)
    for cell,text in zip(table.rows[0].cells,['产品编号','产品名称','品牌','数量']):cell.text=text
    for i in range(24):
        cells=table.add_row().cells
        for cell,value in zip(cells,[f'W1400700{i:04d}','按钮 指示灯','施耐德','17']):cell.text=value
    d.save(path)


def test_cjk_export_chinese_and_tables(tmp_path):
    path=tmp_path/'物料清单.docx';make_docx(path)
    pdf=readable_docx_to_pdf(path,tmp_path/'export.pdf')
    verify_chinese_pdf(pdf,'电气安装材料清单 中文能够正常显示，原始资料需要完整。'+'产品编号产品名称品牌数量'+('按钮 指示灯施耐德')*24)
    r=PdfReader(pdf)
    text='\n'.join(page.extract_text() for page in r.pages)
    assert text.count('施耐德') == 24
    assert text.count('W1400700') == 24
    assert '中文能够正常显示' in text
    assert any('/FontFile2' in str(font.get_object().get('/FontDescriptor').get_object()) for font in r.pages[0]['/Resources']['/Font'].get_object().values() if font.get_object().get('/FontDescriptor'))


def test_unreadable_office_pdf_falls_back_instead_of_blocking(tmp_path):
    src=tmp_path/'清单.docx';make_docx(src)
    # Mac can produce a PDF whose text is not recoverable despite success status.
    def fake_broken_converter(source, folder):
        from reportlab.pdfgen.canvas import Canvas
        bad=folder/'清单.pdf';c=Canvas(str(bad));c.drawString(40,40,'aaa');c.save()
        return bad
    result=convert_with_font_retries(src,tmp_path,fake_broken_converter)
    assert result.name.endswith('_可读重排.pdf')
    assert '施耐德' in ''.join(p.extract_text() for p in PdfReader(result).pages)


def test_ui_does_not_repeat_long_marketing_explanations():
    html=(Path(__file__).parents[1]/'web/index.html').read_text()
    assert '不必先寻找工具' not in html
    assert '所有结果另存为新文件' not in html
    assert '原始文件不会被修改' not in html
    assert '推荐只是建议' not in html
    assert 'V1.0.0' in html
