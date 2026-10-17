"""High-volume procurement workbook and async job regressions."""
from __future__ import annotations
import csv
import os
from io import BytesIO
from pathlib import Path
import time
import zipfile

from docx import Document
from fastapi.testclient import TestClient
import pytest

from filedesk.engine import OPERATIONS, find_libreoffice
from filedesk.server import app
from filedesk.smart_excel import _load_sheets, _sheet_model

REAL = Path(os.environ.get('FILEDESK_QA_XLS','/nonexistent/filedesk-private-qa.xls'))


@pytest.mark.skipif(not REAL.is_file() or not find_libreoffice(),reason='private procurement workbook unavailable')
def test_large_xls_completes_without_giant_word_appendix(tmp_path):
    outputs=OPERATIONS['excel_to_word']([REAL],tmp_path,{})
    docs=[p for p in outputs if p.suffix=='.docx']
    csvs=[p for p in outputs if p.suffix=='.csv']
    assert len(docs)==10
    assert len(csvs)==1
    # Main report is readable/editable, while every source field is exported.
    first=Document(docs[0]);last=Document(docs[-1])
    assert sum(len(t.rows) for t in first.tables)<=430
    assert all('原始 Excel 行号' in ''.join(p.text for p in doc.paragraphs) for doc in (first,last))
    assert any(t.rows[0].cells[0].text=='序号' for t in first.tables)
    with csvs[0].open(encoding='utf-8-sig',newline='') as file:
        reader=csv.reader(file)
        heading=next(reader)
        rows=list(reader)
    assert len(heading)==61  # Source row number + 60 Excel fields.
    assert len(rows)==3612
    assert rows[0][0]=='2'
    assert '采购编号' in heading and '产品名称' in heading
    assert any('45#链轮' in ''.join(row) for row in rows)
    source_sheet=_load_sheets(REAL)[0][1]
    model=_sheet_model('采购列表',source_sheet)
    assert len(rows)==len(model['records'])
    assert rows[0][heading.index('采购编号')]=='CG202502070017'


def test_async_job_reports_progress_downloads_and_has_history(tmp_path,monkeypatch):
    monkeypatch.setenv('FILEDESK_HISTORY_DIR',str(tmp_path/'results'))
    sample=Path(__file__).parent/'fixtures'/'orders.xlsx'
    client=TestClient(app)
    with sample.open('rb') as handle:
        response=client.post('/api/jobs',data={'operation':'excel_to_word','options':'{}'},
                             files=[('files',('orders.xlsx',handle,'application/octet-stream'))])
    assert response.status_code==202,response.text
    info=response.json();identifier=info['id']
    assert len(identifier)==32
    for _ in range(160):
        state=client.get('/api/jobs/'+identifier).json()
        if state['status'] in ('completed','failed'):break
        time.sleep(.1)
    assert state['status']=='completed',state
    assert state['percent']==100
    assert state['result_id']
    result=client.get('/api/jobs/'+identifier+'/download')
    assert result.status_code==200
    assert result.content[:2]==b'PK'  # Editable DOCX
    assert client.get('/api/history/'+state['result_id']+'/download').content==result.content


def test_async_job_surfaces_invalid_workbook_instead_of_hanging(tmp_path,monkeypatch):
    monkeypatch.setenv('FILEDESK_HISTORY_DIR',str(tmp_path/'results'))
    client=TestClient(app)
    response=client.post('/api/jobs',data={'operation':'excel_to_word','options':'{}'},
                         files=[('files',('broken.xlsx',b'invalid xlsx','application/octet-stream'))])
    assert response.status_code==202
    identifier=response.json()['id']
    for _ in range(100):
        state=client.get('/api/jobs/'+identifier).json()
        if state['status'] in ('completed','failed'):break
        time.sleep(.1)
    assert state['status']=='failed'
    assert state['error']
    assert client.get('/api/jobs/'+identifier+'/download').status_code==404
