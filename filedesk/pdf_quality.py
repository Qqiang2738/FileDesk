"""Validate that converted PDFs retain their source Chinese text and embed fonts.

A PDF's BaseFont name is *not* a reliable language/coverage test: LibreOffice
regularly emits subset aliases, especially on macOS.  Validation therefore
uses text fidelity and the PDF's actual embedded font resources instead of
an allow-list of family names. Files never leave the local machine.
"""
from __future__ import annotations

from collections import Counter
from pathlib import Path
import platform
import re
import tempfile

from pypdf import PdfReader
from .engine import FileDeskError

CHINESE = re.compile(r'[\u3400-\u9fff]')
FONT_ERROR = 'PDF 中文字体或内容校验未通过，已中止导出，避免产生乱码。'


def pdf_diagnostics(path: Path) -> dict:
    try:
        reader = PdfReader(str(path), strict=False)
        if not reader.pages:
            raise ValueError('empty pdf')
        texts = []
        fonts = []
        for page in reader.pages:
            texts.append(page.extract_text() or '')
            resources = page.get('/Resources')
            if not resources:
                continue
            font_refs = resources.get_object().get('/Font')
            if not font_refs:
                continue
            for ref in font_refs.get_object().values():
                font = ref.get_object()
                base = str(font.get('/BaseFont', ''))
                candidates = [font]
                if font.get('/DescendantFonts'):
                    candidates.extend(obj.get_object() for obj in font['/DescendantFonts'])
                is_embedded = False
                for candidate in candidates:
                    descriptor = candidate.get('/FontDescriptor')
                    if descriptor:
                        descriptor = descriptor.get_object()
                        is_embedded |= any(descriptor.get(k) for k in ('/FontFile', '/FontFile2', '/FontFile3'))
                fonts.append({'name': base, 'embedded': is_embedded,
                              'unicode': bool(font.get('/ToUnicode'))})
        return {'text': '\n'.join(texts), 'fonts': fonts, 'pages': len(reader.pages)}
    except Exception as exc:
        raise FileDeskError('生成的 PDF 无法正常打开，转换已中止。') from exc


def _font_evidence(details: dict) -> bool:
    # When the output uses meaningful Chinese text it should contain at least
    # one embedded, ToUnicode-mapped font. Do not require a specific font NAME:
    # subset aliases vary between macOS, Windows, and Linux.
    return any(f['embedded'] and f['unicode'] for f in details['fonts'])


def verify_chinese_pdf(path: Path, expected_text: str = '') -> None:
    details = pdf_diagnostics(path)
    actual = details['text']
    # Some macOS LibreOffice font subsets confuse pypdf even if the PDF has
    # correct Unicode mapping. Cross-check with an independent extractor.
    try:
        import pymupdf
        with pymupdf.open(str(path)) as doc:
            alternative='\n'.join(page.get_text(sort=True) for page in doc)
        if alternative:
            def count_match(text):
                wanted=Counter(CHINESE.findall(expected_text))
                found=Counter(CHINESE.findall(text))
                return (sum(min(n,found[ch]) for ch,n in wanted.items()) /
                        max(1,sum(wanted.values())))
            if count_match(alternative)>count_match(actual):
                actual=alternative
    except (ImportError, OSError, RuntimeError, ValueError):
        pass
    if expected_text.strip() and not actual.strip():
        raise FileDeskError('PDF 没有可读取的文字，转换已中止。')
    wanted = Counter(CHINESE.findall(expected_text))
    observed = Counter(CHINESE.findall(actual))
    if len(CHINESE.findall(expected_text)) >= 5:
        if not observed:
            raise FileDeskError(FONT_ERROR + ' PDF 未能保留中文文字。')
        # Detect the "Chinese became the same few characters" failure that
        # a simplistic character count or whitelist missed in earlier builds.
        recall = sum(min(count, observed[char]) for char, count in wanted.items()) / sum(wanted.values())
        if recall < 0.84:
            raise FileDeskError(FONT_ERROR + f' 原文字内容匹配度不足（{recall:.0%}）。')
        if not _font_evidence(details):
            raise FileDeskError(FONT_ERROR + ' PDF 未能确认嵌入可映射字体。')
    elif len(CHINESE.findall(actual)) >= 5 and not _font_evidence(details):
        raise FileDeskError(FONT_ERROR + ' PDF 未能确认嵌入字体。')


def _rewrite_fonts(src: Path, target: Path, font_name: str) -> None:
    """Set direct and inherited fonts in a *temporary* DOCX for an LO retry.

    Include styles/table-style defaults and headers/footers, not only Chinese
    runs. This corrects documents where the East Asian font is inherited.
    """
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(src)
    def set_run_font(rpr):
        if rpr is None:
            return
        fonts = rpr.rFonts
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            rpr.insert(0, fonts)
        for kind in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + kind), font_name)
        for kind in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + kind), None)

    for style in doc.styles:
        set_run_font(style._element.get_or_add_rPr())
    defaults = doc.styles.element.find(qn('w:docDefaults'))
    if defaults is not None:
        rpr_default = defaults.find(qn('w:rPrDefault'))
        if rpr_default is not None:
            default_rpr = rpr_default.find(qn('w:rPr'))
            if default_rpr is None:
                default_rpr = OxmlElement('w:rPr')
                rpr_default.append(default_rpr)
            set_run_font(default_rpr)

    def walk(paragraphs, tables):
        for paragraph in paragraphs:
            for run in paragraph.runs:
                set_run_font(run._element.get_or_add_rPr())
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    walk(cell.paragraphs, cell.tables)
    walk(doc.paragraphs, doc.tables)
    for section in doc.sections:
        for region in (section.header, section.footer, section.first_page_header,
                       section.first_page_footer, section.even_page_header, section.even_page_footer):
            walk(region.paragraphs, region.tables)
    doc.save(target)


def _retry_fonts() -> tuple[str, ...]:
    if platform.system() == 'Darwin':
        # Bundled OS font collections; no font files are shipped in FileDesk.
        return ('PingFang SC', 'Songti SC', 'Heiti SC', 'Arial Unicode MS', 'STHeiti')
    if platform.system() == 'Windows':
        return ('Microsoft YaHei', 'SimSun', 'Microsoft JhengHei')
    return ('Noto Sans CJK SC', 'Noto Serif CJK SC', 'WenQuanYi Micro Hei')


def convert_with_font_retries(src: Path, outdir: Path, converter, expected_text: str = '') -> Path:
    """Prefer original Office layout; on DOCX corruption use a distinct CJK PDF renderer.

    A poor extraction score is NOT repaired by cycling through near-identical
    LibreOffice font substitutions. Reflow is explicit in the resulting filename.
    """
    if not expected_text and src.suffix.lower() == '.docx':
        expected_text = extract_docx_text(src)
    initial_failure = None
    try:
        converted = converter(src, outdir)
        verify_chinese_pdf(converted, expected_text)
        return converted
    except FileDeskError as exc:
        initial_failure=exc
        if src.suffix.lower() != '.docx':
            raise

    # Independent renderer: does not ask LibreOffice to substitute Mac fonts.
    from .readable_pdf import readable_docx_to_pdf
    fallback=outdir/(src.stem+'_可读重排.pdf')
    try:
        readable_docx_to_pdf(src,fallback)
        verify_chinese_pdf(fallback, expected_text)
        return fallback
    except FileDeskError as fallback_error:
        raise FileDeskError(
            f'LibreOffice 生成的 PDF 中文异常；备用可读排版也无法完整保留内容。'
            f'原始错误：{initial_failure}。备用错误：{fallback_error}'
        ) from fallback_error


def extract_docx_text(src: Path) -> str:
    """Read paragraphs and tables in their displayed order for QA."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document=Document(src)
    parts=[]
    for block in document.iter_inner_content():
        if isinstance(block,Paragraph):parts.append(block.text)
        elif isinstance(block,Table):
            for row in block.rows:
                parts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        for region in (section.header,section.footer):
            parts.extend(p.text for p in region.paragraphs)
            for table in region.tables:
                parts.extend(c.text for row in table.rows for c in row.cells)
    return '\n'.join(parts)
