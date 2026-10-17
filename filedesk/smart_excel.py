"""Excel -> structured, editable Word / readable PDF.

Spreadsheet cells are read through the internal OOXML streaming / bounded
reader. Legacy .xls is normalized to .xlsx by isolated local LibreOffice.
CSV is read as delimited text. Original source files are never modified.
"""
from __future__ import annotations

import json
import logging
import csv
import math
import os
import platform
import re
import subprocess
import tempfile
from collections import Counter
from datetime import datetime, date, timedelta
from pathlib import Path
from typing import Any

from .engine import FileDeskError, clean_name, find_libreoffice

COL_LIMIT = 128
ROW_LIMIT = 12000
CELL_LIMIT = 250000
CHUNK_ROWS = 400
# A verbose source-data appendix grows as rows * varying fields. For large
# workbooks keep Word readable and put every field in a separate CSV instead.
LARGE_APPENDIX_CELLS = 6500
SHEET_LIMIT = 8
DARK = '183E33'
GREEN = '2D644F'
GRAY = '64736B'
LINE = 'E5EBE7'
WHITE = 'FFFFFF'


def _col_index(label: str) -> int:
    val = 0
    for ch in label.upper():
        val = val * 26 + ord(ch) - ord('A') + 1
    return val


def _libreoffice_convert(src: Path, output_dir: Path, extension: str) -> Path:
    exe = find_libreoffice()
    if not exe:
        raise FileDeskError('该格式需要本机安装 LibreOffice（免费），安装后再处理。')
    with tempfile.TemporaryDirectory(prefix='filedesk-office-profile-') as profile:
        command = [exe, '-env:UserInstallation=' + Path(profile).as_uri(),
                   '--headless', '--convert-to', extension, '--outdir', str(output_dir), str(src)]
        try:
            result = subprocess.run(command, capture_output=True, text=True, timeout=120)
        except subprocess.TimeoutExpired as exc:
            raise FileDeskError('Office 文件转换超时。建议先在表格软件中确认文件可以正常打开。') from exc
    converted = output_dir / (src.stem + '.' + extension)
    if result.returncode or not converted.is_file() or converted.stat().st_size == 0:
        logging.error('LibreOffice import failure (%s): stdout=%s stderr=%s', result.returncode, result.stdout[-800:], result.stderr[-800:])
        raise FileDeskError(f'LibreOffice 无法读取此 XLS 文件（退出码 {result.returncode}）。请尝试用 LibreOffice 打开并另存为 XLSX 后再导入。')
    return converted


def _load_sheets(source: Path) -> list[tuple[str, list[list[Any]]]]:
    from .xlsx_reader import read_workbook
    with tempfile.TemporaryDirectory(prefix='filedesk-sheet-') as tmp:
        target=source
        if source.suffix.lower()=='.xls':
            target=_libreoffice_convert(source, Path(tmp),'xlsx')
        return read_workbook(target, max_rows=ROW_LIMIT, max_cols=COL_LIMIT,
                             max_cells=CELL_LIMIT,max_sheets=SHEET_LIMIT)


def _blank(v: Any) -> bool:
    return v is None or isinstance(v, str) and not v.strip()


def _display(v: Any, header: str = '') -> str:
    if _blank(v):
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d %H:%M' if v.hour or v.minute else '%Y-%m-%d')
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return '是' if v else '否'
    if isinstance(v, (int, float)):
        if isinstance(v, float) and not math.isfinite(v):
            return str(v)
        if re.search('日期|时间|date', header, re.I) and 1 <= v <= 2958465:
            try:
                return (datetime(1899, 12, 30) + timedelta(days=float(v))).strftime('%Y-%m-%d')
            except OverflowError:
                pass
        return str(int(v)) if float(v).is_integer() else f'{v:.9f}'.rstrip('0').rstrip('.')
    return str(v).strip()


def _excel_col(number: int) -> str:
    result = ''
    while number:
        number,mod=divmod(number-1,26)
        result=chr(65+mod)+result
    return result


def _sheet_model(name: str, matrix: list[list[Any]]) -> dict:
    # Identify a real header row; do not infer unlabeled cells as if they were clean columns.
    best = (-1, -1)
    for idx, row in enumerate(matrix[:12]):
        labels = [_display(v) for v in row if not _blank(v)]
        distinct = len(set(labels))
        named = sum(bool(re.search(r'[\u4e00-\u9fffA-Za-z]', v)) for v in labels)
        score = distinct * 2 + named - sum(len(v)>70 for v in labels)*2
        score -= 5 * sum(isinstance(v, (int, float, datetime, date)) for v in row if not _blank(v))
        score += 4 * sum(v in ('产品编号','产品名称','数量','单位','清单主题','订单编号','型号','名称','日期','客户','产品') for v in labels)
        if named >= 2 and distinct >= 2 and score > best[0]:
            best = (score, idx)
    if best[1] == -1:
        raise FileDeskError(f'工作表「{name}」未识别到可靠的表头，不能安全地智能排版。')
    header_row = best[1]
    # Keep title/notes above a detected header in the source appendix.
    # Their original row/column coordinates remain visible for verification.
    preamble = [(f'{_excel_col(ci + 1)}{ri + 1}', _display(v))
                for ri,row in enumerate(matrix[:header_row])
                for ci,v in enumerate(row) if not _blank(v)]
    raw_headings = matrix[header_row]
    for raw_row in matrix[header_row + 1:]:
        for ci, value in enumerate(raw_row):
            if not _blank(value) and (ci >= len(raw_headings) or _blank(raw_headings[ci])):
                raise FileDeskError(f'工作表「{name}」有未命名列包含数据。为防止丢字段，请先补全表头再转换。')
    valid = [(i, _display(h)) for i,h in enumerate(raw_headings) if not _blank(h)]
    headings = []
    used = Counter()
    for col, label in valid:
        used[label] += 1
        headings.append((col, label if used[label] == 1 else f'{label}（{used[label]}）'))
    if len(headings) > COL_LIMIT:
        raise FileDeskError('有效表头列数超出限制。')
    records = []
    for row_index, row in enumerate(matrix[header_row+1:], header_row+2):
        content = {heading:_display(row[index], heading) if index < len(row) else '' for index,heading in headings}
        if any(content.values()):
            records.append({'excel_row': row_index, 'data':content})
    if not records:
        raise FileDeskError(f'工作表「{name}」只有表头，没有数据。')
    headers = [v for _,v in headings]
    # Only assume a master row where the user supplied a credible hierarchy field.
    master = None
    detail = records
    level_field = next((h for h in headers if h in ('层级','BOM层级','阶层')), None)
    if level_field and len(records)>1 and records[0]['data'].get(level_field) in ('0','0.0'):
        detail = records[1:]
        if all(r['data'].get(level_field,'') not in ('','0','0.0') for r in detail):
            master = records[0]
        else:
            detail = records
    priority = ('产品编号','物料编号','品号','产品名称','物料名称','材料名称','品牌','产品型号','型号','规格','数量','单位')
    selected = [key for key in priority if key in headers]
    # Keep familiar BOM tables compact; only add fallback columns when the
    # recognized set is too sparse to make a useful primary table.
    if len(selected) < 5:
        for h in headers:
            if len(selected) >= 6:
                break
            if h not in selected and h not in ('层级','深度','损耗率','清单主题') and any(r['data'].get(h) for r in detail):
                selected.append(h)
    if len(selected)>7:
        selected=selected[:7]
    if not selected:
        selected = headers[:min(6,len(headers))]
    # '层级' can be shown as row index only when it matches every detail record.
    use_layer = bool(level_field and all(r['data'].get(level_field) == str(i) for i,r in enumerate(detail, 1)))
    other_fields = [h for h in headers if h not in selected and not (use_layer and h==level_field)]
    common=[]
    different=[]
    for h in other_fields:
        values=[r['data'].get(h,'') for r in detail]
        if all(v==values[0] for v in values) and values[0]:
            common.append((h,values[0]))
        elif any(values):
            different.append(h)
    missing=[]
    # Explicit source coverage check: every non-blank spreadsheet field must
    # appear in master / primary table / common attributes / variable appendix.
    for row in records:
        is_master = master is not None and row is master
        for h,value in row['data'].items():
            if not value:continue
            represented = is_master or h in selected or (use_layer and h==level_field) or h in dict(common) or h in different
            if not represented:missing.append((row['excel_row'],h))
    if missing:
        raise FileDeskError('字段校验未通过：智能排版未能保存所有原始字段，已中止输出。')
    return {'sheet':name, 'headers':headers,'header_row':header_row+1,'master':master,'preamble':preamble,
            'detail':detail,'primary':selected,'common':common,'different':different,
            'use_layer':use_layer,'level_field':level_field,'records':records}


def _font_name() -> str:
    if platform.system()=='Darwin':return 'PingFang SC'
    if platform.system()=='Windows':return 'Microsoft YaHei'
    return 'Noto Sans CJK SC'


def _generate_word(models: list[dict], result: Path) -> None:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.section import WD_ORIENTATION
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    font = _font_name()
    document=Document()
    section=document.sections[0]
    section.page_width=Cm(29.7)
    section.page_height=Cm(21.0)
    section.orientation=WD_ORIENTATION.LANDSCAPE
    section.top_margin=Cm(1.30)
    section.bottom_margin=Cm(1.10)
    section.left_margin=Cm(1.3)
    section.right_margin=Cm(1.3)
    section.header_distance=Cm(.65)
    section.footer_distance=Cm(.65)
    normal=document.styles['Normal']
    normal.font.name=font
    normal.font.size=Pt(9)
    normal.paragraph_format.space_after=Pt(4)

    def font_style(run,size=9,color=DARK,bold=False):
        run.bold=bold
        run.font.name=font
        run.font.size=Pt(size)
        run.font.color.rgb=RGBColor.from_string(color)
        rpr=run._element.get_or_add_rPr()
        fonts=rpr.rFonts
        if fonts is None:
            fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
        for attr in ('ascii','hAnsi','eastAsia','cs'):
            fonts.set(qn('w:'+attr),font)

    def shade(cell,color):
        el=OxmlElement('w:shd');el.set(qn('w:fill'),color)
        cell._tc.get_or_add_tcPr().append(el)

    def border(cell):
        b=OxmlElement('w:tcBorders')
        for edge in ('bottom','top','left','right'):
            el=OxmlElement('w:'+edge)
            el.set(qn('w:val'),'single')
            el.set(qn('w:sz'),'4')
            el.set(qn('w:color'),LINE)
            b.append(el)
        cell._tc.get_or_add_tcPr().append(b)

    def no_split(row):
        rowPr=row._tr.get_or_add_trPr()
        rowPr.append(OxmlElement('w:cantSplit'))

    def write_cell(cell,value,*,bold=False,size=8.2,color=DARK,center=False):
        cell.text=''
        p=cell.paragraphs[0]
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.line_spacing=1.12
        if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        font_style(p.add_run(str(value)),size,color,bold)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        border(cell)

    def para(value,*,size=9,color=GRAY,bold=False,after=5):
        p=document.add_paragraph()
        p.paragraph_format.space_after=Pt(after)
        font_style(p.add_run(str(value)),size,color,bold)
        return p

    def pairs_table(items):
        table=document.add_table(rows=0,cols=4)
        table.autofit=False
        widths=[3.2,10.1,3.1,10.1]
        for start in range(0,len(items),2):
            cells=table.add_row().cells
            for position in (0,1):
                k=start+position
                label,value=items[k] if k < len(items) else ('','')
                col=position*2
                cells[col].width=Cm(widths[col]);cells[col+1].width=Cm(widths[col+1])
                write_cell(cells[col],label,bold=True,color=GRAY,size=8)
                shade(cells[col],'F4F7F5')
                write_cell(cells[col+1],value,size=8)
            no_split(table.rows[-1])

    for sheet_idx,model in enumerate(models):
        if sheet_idx:document.add_page_break()
        para('文序  /  FILEDESK',size=9,color=GREEN,bold=True,after=4)
        title = '物料清单' if '产品编号' in model['headers'] and '产品名称' in model['headers'] else '智能排版表格'
        prefix = model['sheet']
        master=model['master']
        if title == '物料清单' and master:
            product = master['data'].get('产品名称','')
            material = re.search(r'[（(]([^（）()]{2,18}材料)[）)]', product)
            if material:
                prefix = material.group(1)
        para(f'{prefix} · {title}',size=18,color=DARK,bold=True,after=5)
        para(('从原始 Excel 直接生成 · 主表可编辑，完整字段随包附送' if model.get('compact_source') else '从原始 Excel 直接生成 · 表格文字可编辑 · 全部非空字段保留'),size=9,after=4)
        if model.get('batch_info'):
            para(model['batch_info'],size=8,color=GREEN,after=7)
        master=model['master']
        short=[]
        if master:
            for k in ('清单编号','产品编号','工程师'):
                if master['data'].get(k):short.append((k,master['data'][k]))
        short.append(('明细数量',f'{len(model["detail"])} 项'))
        if short:pairs_table(short[:4]);para('',after=1)
        para('明细  /  DETAILS',size=10,color=DARK,bold=True,after=6)
        primary=model['primary']
        columns=['序号']+primary
        usable=27.0
        weights=[]
        for h in columns:
            if h=='序号':weights.append(1.25)
            elif any(w in h for w in ('型号','规格','描述','备注','地址','主题')):weights.append(9.6)
            elif any(w in h for w in ('编号','名称')):weights.append(3.2)
            elif any(w in h for w in ('数量','单位','价格')):weights.append(1.7)
            else:weights.append(2.7)
        total=sum(weights)
        widths=[usable*w/total for w in weights]
        t=document.add_table(rows=1,cols=len(columns))
        t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for j,h in enumerate(columns):
            t.columns[j].width=Cm(widths[j]);cell=t.rows[0].cells[j];cell.width=Cm(widths[j])
            shade(cell,DARK)
            write_cell(cell,h,bold=True,size=8,color=WHITE)
        trpr=t.rows[0]._tr.get_or_add_trPr()
        flag=OxmlElement('w:tblHeader');flag.set(qn('w:val'),'true');trpr.append(flag)
        no_split(t.rows[0])
        for n,record in enumerate(model['detail'],1):
            cells=t.add_row().cells
            row_number=(record['data'].get(model['level_field']) if model['use_layer'] and model['level_field'] else None) or str(record['excel_row'])
            vals=[str(row_number)]+[record['data'].get(h,'') for h in primary]
            for j,value in enumerate(vals):
                cells[j].width=Cm(widths[j])
                write_cell(cells[j],value,size=8.4 if j else 8.0,center=(j==0 or columns[j] in ('数量','单位','层级')))
                if n%2==0:shade(cells[j],'F6F9F7')
            no_split(t.rows[-1])
        para('本表直接读取 Excel 数据，行号可与完整字段 CSV 对照。' if model.get('compact_source') else '说明：本表来自 Excel 单元格数据，不经 PDF 文字识别。原始文件保持不变。',size=8,after=0)
        if model.get('compact_source'):
            para('完整来源字段请见压缩包内的 CSV 文件；主表第一列对应原始 Excel 行号。',size=8,after=0)
            continue
        document.add_page_break()
        para('来源数据  /  SOURCE RECORD',size=16,color=DARK,bold=True,after=5)
        para('为方便核对，保留源表中未在明细表展开的字段。',size=9,after=8)
        para(f'来源工作表：{model["sheet"]} · 表头第 {model["header_row"]} 行 · 共 {len(model["records"])} 条记录',size=8,after=7)
        if model['preamble']:
            para('表头之前的原始说明',size=10,color=DARK,bold=True,after=5)
            pairs_table(model['preamble'])
            para('',after=2)
        if master:
            para(f'主记录（Excel 第 {master["excel_row"]} 行）',size=10,color=DARK,bold=True,after=5)
            pairs_table([(k,v) for k,v in master['data'].items() if v])
            para('',after=2)
        if model['common']:
            para('明细公共字段',size=10,color=DARK,bold=True,after=5)
            pairs_table(model['common'])
            para('',after=2)
        if model['different']:
            para('明细差异字段（原始行号对应 Excel 行号）',size=10,color=DARK,bold=True,after=6)
            diff=document.add_table(rows=1,cols=3)
            diff.autofit=False
            for j,name in enumerate(['Excel 行号','字段','原始值']):
                c=diff.rows[0].cells[j];shade(c,DARK);write_cell(c,name,bold=True,color=WHITE)
            trpr=diff.rows[0]._tr.get_or_add_trPr()
            flag=OxmlElement('w:tblHeader');flag.set(qn('w:val'),'true');trpr.append(flag)
            for record in model['detail']:
                for field in model['different']:
                    value=record['data'].get(field,'')
                    if not value:continue
                    cells=diff.add_row().cells
                    for cell,val in zip(cells,(record['excel_row'],field,value)):
                        write_cell(cell,val,size=8)
                    no_split(diff.rows[-1])
        if model['use_layer']:
            para('字段对应：主表“序号”与原始 Excel 的“层级”一致。',size=8,after=0)
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font_style(footer.add_run('FILEDESK  ·  本地智能排版  |  可编辑 Word 表格'),7,GRAY)
    document.save(result)


def _split_models(models: list[dict]) -> list[tuple[list[dict], str]]:
    """Keep normal reports unchanged; large sheets become safe, usable sections.

    All original Excel row numbers remain present in record metadata. Never
    drop columns just to squeeze an oversized table onto a page.
    """
    chunks=[]
    for model in models:
        records=model['detail']
        if len(records)<=CHUNK_ROWS:
            chunks.append(([model], ''))
            continue
        total=math.ceil(len(records)/CHUNK_ROWS)
        for i in range(total):
            piece=records[i*CHUNK_ROWS:(i+1)*CHUNK_ROWS]
            part=dict(model)
            part['detail']=piece
            part['records']=([model['master']] if model['master'] else [])+piece
            part['batch_info']=f'第 {i+1}/{total} 部分 · 原始明细 {len(records)} 行'
            chunks.append(([part], f'_{clean_name(model["sheet"])}_第{i+1:03d}部分'))
    if len(chunks)>100:
        raise FileDeskError('本次智能排版需要超过 100 个分段文件，建议缩小处理范围。')
    return chunks


def _export_full_source_csv(model: dict, target: Path) -> None:
    """Export every modeled cell with original row positions; UTF-8 BOM for Excel.

    The report may show a compact subset. The CSV is the lossless *displayed
    cell-value* companion, not an exact reproduction of Excel styles/formulas.
    Cells starting with formula tokens are prefixed for spreadsheet safety.
    """
    with target.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(['原始 Excel 行号', *model['headers']])
        for record in model['records']:
            def safe(value):
                value = str(value)
                return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value
            writer.writerow([record['excel_row'], *[safe(record['data'].get(h,'')) for h in model['headers']]])


def _convert_streaming(src: Path, out: Path, options: dict, *, pdf: bool, index: int) -> list[Path]:
    """Incremental large/CSV conversion. Never materialize all worksheet cells.

    Each output DOCX has at most CHUNK_ROWS records; all original displayed
    values are exported as UTF-8-SIG CSV with source row numbers. CSV is not a
    byte-identical backup of formulas / formatting, so preserve the original.
    """
    from itertools import chain
    from .tabular_stream import sheets as read_stream, MAX_CELLS
    outputs=[]
    progress=options.get('_progress')
    with tempfile.TemporaryDirectory(prefix='filedesk-stream-') as tmp:
        target=src
        if src.suffix.lower()=='.xls':
            target=_libreoffice_convert(src, Path(tmp), 'xlsx')
        with read_stream(target) as streams:
            for sheet_index,(sheet,rows) in enumerate(streams,1):
                iterator=iter(rows)
                warm=[]
                # Enough to infer the header and normal short-sheet layout.
                for row in iterator:
                    warm.append(row)
                    if len(warm)>=min(CHUNK_ROWS+12,420):
                        break
                if not warm:
                    continue
                # First 12 physical rows determine the header. Do not insert a
                # 50000-row sparse matrix just because row numbers have gaps.
                max_preview_row=min(12,max(n for n,_ in warm))
                candidate=[]
                for i in range(1,max_preview_row+1):
                    record=next((value for number,value in warm if number==i), {})
                    candidate.append([record.get(j) for j in range(max(record.keys(),default=-1)+1)])
                # Rank header candidates using the same scoring as the normal model.
                best=(-1,-1)
                for idx,line in enumerate(candidate):
                    labels=[_display(v) for v in line if not _blank(v)]
                    distinct=len(set(labels))
                    named=sum(bool(re.search(r'[\u4e00-\u9fffA-Za-z]',v)) for v in labels)
                    score=distinct*2+named-sum(len(v)>70 for v in labels)*2
                    score-=5*sum(isinstance(v,(int,float,datetime,date)) for v in line if not _blank(v))
                    score+=4*sum(v in ('产品编号','产品名称','数量','单位','清单主题','订单编号','型号','名称','日期','客户','产品') for v in labels)
                    if named>=2 and distinct>=2 and score>best[0]:
                        best=(score,idx)
                if best[1]<0:
                    raise FileDeskError(f'工作表「{sheet}」未识别到表头。请先在原文件中补全字段名称。')
                header_row=best[1]+1
                heading_cells=next((v for n,v in warm if n==header_row), {})
                positions=sorted(j for j,v in heading_cells.items() if not _blank(v))
                if not positions:
                    raise FileDeskError(f'工作表「{sheet}」缺少表头。')
                headings=[]
                used=Counter()
                for j in positions:
                    value=_display(heading_cells[j])
                    used[value]+=1
                    headings.append(value if used[value]==1 else f'{value}（{used[value]}）')
                fields=dict(zip(positions,headings))
                candidate_rows=[(n,r) for n,r in warm if n>header_row]
                if not candidate_rows and not any(n>header_row for n,_ in warm):
                    # Future rows may be present; do not reject until exhausted.
                    pass
                name=f'{clean_name(src.name)}_{index:02d}_智能排版_{clean_name(sheet)}'
                csv_file=out/(name+'_完整字段.csv')
                count=0;file_count=0;filled=0
                current=[]
                prior_master=None
                primary=None
                def flush():
                    nonlocal file_count, primary
                    if not current:return
                    if file_count>=125:
                        raise FileDeskError('大表格分册超过 125 份；请按业务范围拆分后再处理。')
                    file_count+=1
                    # Infer a useful small number of display columns once.
                    if primary is None:
                        candidates=('产品编号','物料编号','品号','产品名称','物料名称','材料名称','品牌','产品型号','型号','规格','数量','单位')
                        primary=[v for v in candidates if v in headings][:7]
                        if len(primary)<5:
                            for h in headings:
                                if h not in primary and len(primary)<6:
                                    primary.append(h)
                    model={'sheet':sheet,'headers':headings,'header_row':header_row,
                           'master':prior_master,'preamble':[],'detail':current[:],
                           'primary':primary,'common':[],'different':[], 'use_layer':False,
                           'level_field':None,'records':current[:], 'compact_source':True,
                           'batch_info':f'第 {file_count} 部分 · 源表行号 {current[0]["excel_row"]}–{current[-1]["excel_row"]}'}
                    stem=name+f'_第{file_count:03d}部分'
                    if callable(progress):
                        progress(f'正在生成第 {file_count} 份文档 · 已读取 {count:,} 行',min(90,20+file_count))
                    if pdf:
                        from .readable_pdf import readable_docx_to_pdf
                        from .pdf_quality import verify_chinese_pdf,extract_docx_text
                        with tempfile.TemporaryDirectory(prefix='filedesk-cjk-') as doc_temp:
                            docx=Path(doc_temp)/(stem+'.docx')
                            _generate_word([model],docx)
                            dest=out/(stem+'.pdf')
                            readable_docx_to_pdf(docx,dest)
                            verify_chinese_pdf(dest,extract_docx_text(docx))
                            outputs.append(dest)
                    else:
                        dest=out/(stem+'.docx')
                        _generate_word([model],dest)
                        outputs.append(dest)
                    current.clear()
                with csv_file.open('w',encoding='utf-8-sig',newline='') as io:
                    writer=csv.writer(io)
                    writer.writerow(['原始行号',*headings])
                    for line,data in chain(warm,iterator):
                        if line<=header_row:
                            continue
                        count+=1
                        if count>50000:
                            raise FileDeskError('大表格单表超过 50,000 行，请拆分工作表后转换。')
                        if any(j not in fields and not _blank(value) for j,value in data.items()):
                            raise FileDeskError(f'第 {line} 行存在没有表头的字段，已停止以避免丢失数据。')
                        filled+=sum(not _blank(v) for v in data.values())
                        if filled>MAX_CELLS:
                            raise FileDeskError(f'单表超过 {MAX_CELLS:,} 个有效单元格的流式安全上限，请按业务范围拆分。')
                        record={h:_display(data.get(j),h) for j,h in fields.items()}
                        def safe_csv(v):
                            v=str(v)
                            return "'"+v if v.lstrip().startswith(('=','+','-','@')) else v
                        writer.writerow([line,*[safe_csv(record[h]) for h in headings]])
                        if any(record.values()):
                            current.append({'excel_row':line,'data':record})
                            if len(current)>=CHUNK_ROWS:
                                flush()
                flush()
                if not count:
                    csv_file.unlink(missing_ok=True)
                    raise FileDeskError(f'工作表「{sheet}」未找到表格内容。')
                outputs.insert(0,csv_file)
                if callable(progress):
                    progress(f'工作表「{sheet}」已处理 {count:,} 行、生成 {file_count} 份文档',90)
    return outputs


def table_to_csv(files: list[Path],out: Path,opts: dict) -> list[Path]:
    """CSV export for XLS/XLSX/CSV; keep every nonempty column and source row."""
    from .tabular_stream import sheets as stream_sheets, MAX_CELLS
    outputs=[]
    for index,src in enumerate(files,1):
        if src.suffix.lower() not in ('.xls','.xlsx','.csv'):
            raise FileDeskError('CSV 导出支持 XLS、XLSX、CSV。')
        with tempfile.TemporaryDirectory(prefix='filedesk-export-') as tmp:
            target=_libreoffice_convert(src,Path(tmp),'xlsx') if src.suffix.lower()=='.xls' else src
            with stream_sheets(target) as streams:
                for sheet_index,(name,rows) in enumerate(streams,1):
                    output=out/f'{clean_name(src.name)}_{index:02d}_{clean_name(name)}.csv'
                    count=0;cells=0
                    # Export raw displayed content without changing column placement.
                    with output.open('w',encoding='utf-8-sig',newline='') as f:
                        writer=csv.writer(f)
                        for ri, data in rows:
                            count+=1;cells+=len(data)
                            if cells>MAX_CELLS:raise FileDeskError('CSV 导出超过本机单表安全上限。')
                            def safe(value):
                                value=_display(value)
                                return "'"+value if value.lstrip().startswith(('=','+','-','@')) else value
                            writer.writerow([ri]+[safe(data.get(ci,'')) for ci in range(max(data,default=-1)+1)])
                    if count:outputs.append(output)
                    else:output.unlink(missing_ok=True)
    if not outputs:raise FileDeskError('没有可导出的表格内容。')
    return outputs


def convert_excel(files: list[Path], out: Path, options: dict, *, pdf: bool) -> list[Path]:
    if not files or len(files)>20:
        raise FileDeskError('智能排版每次支持 1–20 个 Excel 文件。')
    outputs=[]
    for index,src in enumerate(files,1):
        if src.suffix.lower() not in ('.xls','.xlsx','.csv'):
            raise FileDeskError('支持 .xls、.xlsx 和 .csv，请选择表格文件。')
        if src.suffix.lower() == '.csv' or src.stat().st_size > 7 * 1024 * 1024:
            # Skip the legacy in-memory reader for large compressed workbooks.
            outputs.extend(_convert_streaming(src, out, options, pdf=pdf, index=index))
            continue
        try:
            sheets=_load_sheets(src)
        except FileDeskError as exc:
            if '有效单元格超过本机安全上限' in str(exc) or '数据范围超出本机处理上限' in str(exc) or '行列跨度过大' in str(exc) or 'Excel 文件结构过大' in str(exc):
                outputs.extend(_convert_streaming(src, out, options, pdf=pdf, index=index))
                continue
            raise
        progress = options.get('_progress')
        if callable(progress):
            progress('已读取工作表，正在分析数据结构',18)
        models=[_sheet_model(name, data) for name,data in sheets]
        large=[]
        for model in models:
            appendix_cells=sum(bool(row['data'].get(h)) for row in model['detail'] for h in model['different'])
            if appendix_cells > LARGE_APPENDIX_CELLS:
                model['compact_source']=True
                large.append(model)
        name=f'{clean_name(src.name)}_{index:02d}_智能排版'
        for model in large:
            csv_output=out/(name+'_'+clean_name(model['sheet'])+'_完整字段.csv')
            _export_full_source_csv(model,csv_output)
            outputs.append(csv_output)
        chunks = _split_models(models)
        for segment_index, (segment, suffix) in enumerate(chunks):
            if callable(progress):
                progress(f'正在生成第 {segment_index+1}/{len(chunks)} 份可编辑文档', 24+int(65*segment_index/max(1,len(chunks))))
            stem = name + suffix
            if pdf:
                with tempfile.TemporaryDirectory(prefix='filedesk-smart-pdf-') as directory:
                    work_dir=Path(directory)
                    docx=work_dir/(stem+'.docx')
                    _generate_word(segment, docx)
                    # No LibreOffice PDF stage: export the already-structured
                    # DOCX directly with an embedded Chinese-capable font.
                    from .readable_pdf import readable_docx_to_pdf
                    from .pdf_quality import verify_chinese_pdf, extract_docx_text
                    final=out/(stem+'.pdf')
                    readable_docx_to_pdf(docx,final)
                    verify_chinese_pdf(final, extract_docx_text(docx))
                    outputs.append(final)
            else:
                docx=out/(stem+'.docx')
                _generate_word(segment, docx)
                outputs.append(docx)
        if callable(progress):
            progress('正在整理转换结果',92)
    return outputs


def excel_to_word(files: list[Path], out: Path, opts: dict) -> list[Path]:
    return convert_excel(files, out, opts, pdf=False)


def excel_smart_pdf(files: list[Path], out: Path, opts: dict) -> list[Path]:
    return convert_excel(files, out, opts, pdf=True)
