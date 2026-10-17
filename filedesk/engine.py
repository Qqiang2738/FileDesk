"""Local-only file operations. Inputs are temporary, never sent to third-party services."""
from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageOps, UnidentifiedImageError
from pypdf import PdfReader, PdfWriter

PDF_TYPES = {'.pdf'}
IMAGE_TYPES = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff', '.gif'}
OFFICE_TYPES = {'.doc', '.docx', '.odt', '.rtf', '.xls', '.xlsx', '.ods', '.ppt', '.pptx', '.odp'}
MAX_PAGES = 1200


class FileDeskError(Exception):
    pass


def clean_name(name: str) -> str:
    stem = Path(name.replace('\\', '/').split('/')[-1]).stem
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', stem).strip(' .')
    return (stem or '文件')[:90]


def _check_count(paths: list[Path], min_count: int, max_count: int | None = None):
    if len(paths) < min_count or (max_count and len(paths) > max_count):
        if min_count == max_count:
            raise FileDeskError(f'该操作需要 {min_count} 个文件')
        raise FileDeskError(f'该操作至少需要 {min_count} 个文件')


def _check_ext(paths: list[Path], allowed: set[str]):
    for path in paths:
        if path.suffix.lower() not in allowed:
            raise FileDeskError(f'文件类型不支持：{path.name}')


def _reader(path: Path) -> PdfReader:
    try:
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            try:
                if not reader.decrypt(''):
                    raise FileDeskError(f'{path.name} 已加密，请先解锁')
            except FileDeskError:
                raise
            except Exception as e:
                raise FileDeskError(f'{path.name} 已加密，请先解锁') from e
        if len(reader.pages) > MAX_PAGES:
            raise FileDeskError(f'当前限制单个 PDF 最多 {MAX_PAGES} 页')
        return reader
    except FileDeskError:
        raise
    except Exception as e:
        raise FileDeskError(f'无法读取 PDF：{path.name}（可能已损坏）') from e


def parse_pages(spec: str, total: int, *, require: bool = True) -> list[int]:
    """Convert 1-based page selectors into unique, ordered 0-based page indices."""
    if not spec.strip():
        if require:
            raise FileDeskError('请输入页码，例如 1-3,5,8')
        return list(range(total))
    result = []
    found = set()
    for token in re.split(r'[,，;；\s]+', spec.strip()):
        if not token:
            continue
        hit = re.fullmatch(r'(\d+)(?:\s*-\s*(\d+))?', token)
        if not hit:
            raise FileDeskError(f'无效页码：{token}，可用格式如 1-3,5,8')
        start, stop = int(hit.group(1)), int(hit.group(2) or hit.group(1))
        if not (1 <= start <= stop <= total):
            raise FileDeskError(f'页码 {token} 超出有效范围（1-{total}）')
        for p in range(start - 1, stop):
            if p not in found:
                result.append(p)
                found.add(p)
    if not result:
        raise FileDeskError('没有选中任何页面')
    return result


def _write_pdf(writer: PdfWriter, out: Path):
    with out.open('wb') as fp:
        writer.write(fp)
    return [out]


def pdf_merge(files: list[Path], out: Path, opts: dict) -> list[Path]:
    _check_count(files, 2)
    _check_ext(files, PDF_TYPES)
    writer = PdfWriter()
    for path in files:
        writer.append(_reader(path), import_outline=False)
    return _write_pdf(writer, out / '合并后的文件.pdf')


def pdf_split(files: list[Path], out: Path, opts: dict) -> list[Path]:
    _check_count(files, 1, 1)
    _check_ext(files, PDF_TYPES)
    reader = _reader(files[0])
    if len(reader.pages) > 300:
        raise FileDeskError('单页拆分最多支持 300 页，请先提取页码范围')
    results = []
    for i, page in enumerate(reader.pages, 1):
        writer = PdfWriter()
        writer.add_page(page)
        target = out / f'{clean_name(files[0].name)}_第{i:03d}页.pdf'
        _write_pdf(writer, target)
        results.append(target)
    return results


def pdf_extract(files: list[Path], out: Path, opts: dict) -> list[Path]:
    _check_count(files, 1, 1)
    _check_ext(files, PDF_TYPES)
    reader = _reader(files[0])
    pages = parse_pages(str(opts.get('pages', '')), len(reader.pages))
    writer = PdfWriter()
    for p in pages:
        writer.add_page(reader.pages[p])
    return _write_pdf(writer, out / (clean_name(files[0].name) + '_提取.pdf'))


def pdf_remove(files: list[Path], out: Path, opts: dict) -> list[Path]:
    _check_count(files, 1, 1)
    _check_ext(files, PDF_TYPES)
    reader = _reader(files[0])
    removed = set(parse_pages(str(opts.get('pages', '')), len(reader.pages)))
    if len(removed) == len(reader.pages):
        raise FileDeskError('不能删除全部页面')
    writer = PdfWriter()
    for i, page in enumerate(reader.pages):
        if i not in removed:
            writer.add_page(page)
    return _write_pdf(writer, out / (clean_name(files[0].name) + '_已删除页面.pdf'))


def pdf_rotate(files: list[Path], out: Path, opts: dict) -> list[Path]:
    _check_count(files, 1, 1)
    _check_ext(files, PDF_TYPES)
    angle = int(opts.get('angle', 90))
    if angle not in (90, 180, 270):
        raise FileDeskError('旋转角度只能为 90、180 或 270 度')
    reader = _reader(files[0])
    pages = set(parse_pages(str(opts.get('pages', '')), len(reader.pages), require=False))
    writer = PdfWriter()
    for i, page in enumerate(reader.pages):
        if i in pages:
            page.rotate(angle)
        writer.add_page(page)
    return _write_pdf(writer, out / (clean_name(files[0].name) + '_已旋转.pdf'))


def pdf_to_images(files: list[Path], out: Path, opts: dict) -> list[Path]:
    _check_count(files, 1, 1)
    _check_ext(files, PDF_TYPES)
    import pypdfium2 as pdfium
    fmt = str(opts.get('format', 'png')).lower()
    if fmt not in ('png', 'jpg'):
        raise FileDeskError('仅支持 PNG 或 JPG')
    dpi = int(opts.get('dpi', 150))
    if dpi not in (96, 150, 200):
        raise FileDeskError('不支持此分辨率')
    try:
        pdf = pdfium.PdfDocument(str(files[0]))
        if len(pdf) > MAX_PAGES:
            raise FileDeskError('PDF 页数超过处理上限')
        result = []
        for i in range(len(pdf)):
            page = pdf[i]
            bitmap = page.render(scale=dpi / 72)
            image = bitmap.to_pil()
            target = out / f'{clean_name(files[0].name)}_第{i+1:03d}页.{fmt}'
            if fmt == 'jpg':
                image = image.convert('RGB')
                image.save(target, 'JPEG', quality=90, optimize=True)
            else:
                image.save(target, 'PNG', optimize=True)
            result.append(target)
            image.close()
            page.close()
        pdf.close()
        return result
    except FileDeskError:
        raise
    except Exception as e:
        raise FileDeskError('PDF 页面渲染失败，文件可能已损坏或已加密') from e


def _load_image(path: Path) -> Image.Image:
    try:
        with Image.open(path) as original:
            im = ImageOps.exif_transpose(original)
            im.seek(0)
            im.load()
            return im.copy()
    except (OSError, UnidentifiedImageError, ValueError, Image.DecompressionBombError) as e:
        raise FileDeskError(f'无法读取图片：{path.name}') from e


def _flatten_white(img: Image.Image) -> Image.Image:
    if img.mode in ('RGBA', 'LA') or 'transparency' in img.info:
        rgba = img.convert('RGBA')
        base = Image.new('RGB', rgba.size, (255, 255, 255))
        base.paste(rgba, mask=rgba.getchannel('A'))
        return base
    return img.convert('RGB')


def images_to_pdf(files: list[Path], out: Path, opts: dict) -> list[Path]:
    _check_count(files, 1)
    _check_ext(files, IMAGE_TYPES)
    rgb_images = []
    try:
        for path in files:
            image = _load_image(path)
            rgb_images.append(_flatten_white(image))
            image.close()
        target = out / '图片合成.pdf'
        rgb_images[0].save(target, 'PDF', resolution=100.0, save_all=True,
                           append_images=rgb_images[1:])
        return [target]
    finally:
        for image in rgb_images:
            image.close()


def image_convert(files: list[Path], out: Path, opts: dict) -> list[Path]:
    _check_count(files, 1)
    _check_ext(files, IMAGE_TYPES)
    fmt = str(opts.get('format', 'png')).lower()
    if fmt not in ('png', 'jpg', 'webp'):
        raise FileDeskError('请选择 PNG、JPG 或 WebP')
    quality = int(opts.get('quality', 85))
    if not 30 <= quality <= 100:
        raise FileDeskError('质量只能为 30-100')
    max_side = int(opts.get('max_side', 0))
    if max_side not in (0, 1280, 1920, 2560):
        raise FileDeskError('不支持该图片尺寸选项')
    results = []
    for i, path in enumerate(files, 1):
        img = _load_image(path)
        try:
            if max_side:
                img.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
            target = out / f'{clean_name(path.name)}_{i:02d}.{fmt}'
            if fmt == 'jpg':
                _flatten_white(img).save(target, 'JPEG', quality=quality, optimize=True)
            elif fmt == 'webp':
                img.convert('RGBA' if 'A' in img.getbands() else 'RGB').save(
                    target, 'WEBP', quality=quality, method=5)
            else:
                img.convert('RGBA' if 'A' in img.getbands() else 'RGB').save(
                    target, 'PNG', optimize=True)
            results.append(target)
        finally:
            img.close()
    return results


def pdf_to_word(files: list[Path], out: Path, opts: dict) -> list[Path]:
    """Three modes: geometry-aware editable tables, text-only, page images."""
    _check_count(files, 1, 1)
    _check_ext(files, PDF_TYPES)
    from docx import Document
    from docx.shared import Cm
    mode = str(opts.get('mode', 'structured'))
    if mode not in ('structured', 'text', 'visual'):
        raise FileDeskError('请选择智能表格、文本可编辑或页面保真模式')
    if mode == 'structured':
        from .pdf_word import pdf_tables_to_docx
        return [pdf_tables_to_docx(files[0], out / (clean_name(files[0].name) + '_可编辑表格.docx'))]
    filename = out / (clean_name(files[0].name) + '_转换.docx')
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    if mode == 'text':
        reader = _reader(files[0])
        any_text = False
        for i, page in enumerate(reader.pages):
            if i:
                doc.add_page_break()
            text = (page.extract_text(extraction_mode='layout') or '').strip()
            if text:
                any_text = True
                for line in text.splitlines():
                    doc.add_paragraph(line)
        if not any_text:
            raise FileDeskError('没有识别到可编辑文字。这可能是扫描件，请使用版式保真模式；OCR 将在后续版本提供。')
    else:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(files[0]))
        if len(pdf) > 150:
            raise FileDeskError('版式保真模式最多处理 150 页')
        for i in range(len(pdf)):
            if i:
                doc.add_page_break()
            page = pdf[i]
            bitmap = page.render(scale=1.45)
            img = bitmap.to_pil().convert('RGB')
            buffer = io.BytesIO()
            img.save(buffer, 'JPEG', quality=84, optimize=True)
            buffer.seek(0)
            doc.add_picture(buffer, width=Cm(17.4))
            page.close()
            img.close()
        pdf.close()
    doc.save(filename)
    return [filename]


def find_libreoffice() -> str | None:
    for binary in ('soffice', 'libreoffice'):
        path = shutil.which(binary)
        if path:
            return path
    for fixed in ('/Applications/LibreOffice.app/Contents/MacOS/soffice',
                  'C:/Program Files/LibreOffice/program/soffice.exe',
                  'C:/Program Files (x86)/LibreOffice/program/soffice.exe'):
        if Path(fixed).is_file():
            return fixed
    return None


def office_to_pdf(files: list[Path], out: Path, opts: dict) -> list[Path]:
    """Default .docx to clear CJK reflow; other Office formats use LibreOffice.

    The direct renderer preserves editable text/table content, not pixel-identical
    Word pagination. Selecting 'original' explicitly requests LO first.
    """
    _check_count(files, 1)
    _check_ext(files, OFFICE_TYPES)
    mode=str(opts.get('mode','readable'))
    if mode not in ('readable','original'):
        raise FileDeskError('Word 转 PDF 只支持清晰排版或原版式模式')
    exe=find_libreoffice()
    from .pdf_quality import convert_with_font_retries,verify_chinese_pdf,extract_docx_text
    from .readable_pdf import readable_docx_to_pdf
    outputs=[]

    def run_conversion(src: Path, dest: Path) -> Path:
        if not exe:
            raise FileDeskError('此格式或原版式转换需要安装 LibreOffice。')
        dest.mkdir(parents=True,exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='filedesk-lo-') as profile:
            command=[exe,'-env:UserInstallation='+Path(profile).as_uri(),
                     '--headless','--convert-to','pdf','--outdir',str(dest),str(src)]
            try:
                completed=subprocess.run(command,capture_output=True,text=True,timeout=180)
            except subprocess.TimeoutExpired as exc:
                raise FileDeskError('Office 转 PDF 超时，请检查文件大小或是否损坏') from exc
        pdfs=list(dest.glob('*.pdf'))
        if completed.returncode!=0 or not pdfs:
            raise FileDeskError(f'{src.name} 无法转换为 PDF，请检查 LibreOffice 能否打开它')
        return max(pdfs,key=lambda f:f.stat().st_mtime)

    for i,src in enumerate(files,1):
        if src.suffix.lower()=='.docx' and mode=='readable':
            target=out/f'{clean_name(src.name)}_{i:02d}_清晰排版.pdf'
            readable_docx_to_pdf(src,target)
            verify_chinese_pdf(target,extract_docx_text(src))
            outputs.append(target)
            continue
        target_dir=out/f'office_{i}'
        pdf=convert_with_font_retries(src,target_dir,run_conversion,
                extract_docx_text(src) if src.suffix.lower()=='.docx' else '')
        marker='_可读重排' if pdf.stem.endswith('_可读重排') else ''
        final=out/f'{clean_name(src.name)}_{i:02d}{marker}.pdf'
        shutil.copyfile(pdf,final)
        outputs.append(final)
    return outputs


def pdf_compress(files: list[Path], out: Path, opts: dict) -> list[Path]:
    _check_count(files, 1, 1)
    _check_ext(files, PDF_TYPES)
    mode = str(opts.get('mode', 'lossless'))
    target = out / (clean_name(files[0].name) + '_优化.pdf')
    if mode == 'smaller':
        gs = shutil.which('gs') or shutil.which('gswin64c') or next((p for p in ('/opt/homebrew/bin/gs', '/usr/local/bin/gs') if Path(p).is_file()), None)
        if not gs:
            raise FileDeskError('此压缩模式需要本机安装 Ghostscript；也可以选择无损优化')
        command = [gs, '-sDEVICE=pdfwrite', '-dCompatibilityLevel=1.6',
                   '-dPDFSETTINGS=/ebook', '-dDetectDuplicateImages=true',
                   '-dNOPAUSE', '-dBATCH', '-dQUIET', '-sOutputFile=' + str(target),
                   str(files[0])]
        try:
            run = subprocess.run(command, capture_output=True, text=True, timeout=120)
            if run.returncode != 0 or not target.exists():
                raise FileDeskError('压缩失败，请检查 PDF 文件是否正常')
        except subprocess.TimeoutExpired as e:
            raise FileDeskError('PDF 压缩超时') from e
    elif mode == 'lossless':
        reader = _reader(files[0])
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        for page in writer.pages:
            try:
                page.compress_content_streams()
            except Exception:
                pass
        writer.compress_identical_objects(remove_identicals=True, remove_orphans=True)
        _write_pdf(writer, target)
    else:
        raise FileDeskError('未知压缩模式')
    # Never deliver a larger result as a supposed compression win.
    if target.stat().st_size >= files[0].stat().st_size:
        shutil.copy2(files[0], target)
    return [target]


# Lazy import keeps shared FileDeskError and LibreOffice discovery independent.
def excel_to_word(files: list[Path], out: Path, opts: dict) -> list[Path]:
    from .smart_excel import excel_to_word as impl
    return impl(files, out, opts)

def excel_smart_pdf(files: list[Path], out: Path, opts: dict) -> list[Path]:
    from .smart_excel import excel_smart_pdf as impl
    return impl(files, out, opts)

OPERATIONS = {
    'excel_to_word': excel_to_word,
    'excel_smart_pdf': excel_smart_pdf,
    'pdf_merge': pdf_merge,
    'pdf_split': pdf_split,
    'pdf_extract': pdf_extract,
    'pdf_remove': pdf_remove,
    'pdf_rotate': pdf_rotate,
    'pdf_to_images': pdf_to_images,
    'images_to_pdf': images_to_pdf,
    'image_convert': image_convert,
    'table_to_csv': lambda files, out, opts: __import__('filedesk.smart_excel',fromlist=['table_to_csv']).table_to_csv(files,out,opts),
    'pdf_to_word': pdf_to_word,
    'office_to_pdf': office_to_pdf,
    'pdf_compress': pdf_compress,
}
